from __future__ import annotations

import io
import os
from typing import Optional

from fastapi import HTTPException
from PIL import Image

from ...core.config import settings


MB = 1024 * 1024
MAX_UPLOAD_BYTES = settings.CHAT_MAX_UPLOAD_MB * MB

# PIL decompression-bomb protection: 50 MP is generous for any business document.
Image.MAX_IMAGE_PIXELS = 50_000_000

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".csv", ".json", ".docx", ".xlsx", ".pptx", ".png", ".jpg", ".jpeg"}
IMAGE_EXTENSIONS   = {".png", ".jpg", ".jpeg"}
OFFICE_EXTENSIONS  = {".docx", ".xlsx", ".pptx"}
TEXT_EXTENSIONS    = {".txt", ".csv", ".json"}

DIRECT_MIME = {
    ".pdf":  "application/pdf",
    ".txt":  "text/plain",
    ".csv":  "text/csv",
    ".json": "application/json",
    ".png":  "image/png",
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
}


def get_ext(filename: str) -> str:
    _, ext = os.path.splitext(filename.lower())
    return ext


def truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[Content truncated at extraction limit.]"


def extract_text_blocks(response) -> str:
    blocks = [b.text for b in response.content if getattr(b, "type", None) == "text"]
    return "\n".join(blocks).strip() or "No text content in response."


def prepare_upload(prompt: str, safe_name: str, ext: str, data: bytes) -> tuple[str, str, bytes, dict]:
    """
    Return (filename, mime_type, bytes, content_block_template) for the Anthropic Files API.
    The caller fills `source.file_id` in the template after the upload.
    """
    if ext in IMAGE_EXTENSIONS:
        _validate_image(data)
        block = {"type": "image", "source": {"type": "file", "file_id": ""}}
        return safe_name, DIRECT_MIME[ext], data, block

    if ext in OFFICE_EXTENSIONS:
        extracted = _extract_office_text(ext, data)
        if not extracted.strip():
            raise HTTPException(status_code=400, detail="No readable text found in file.")
        extracted = truncate(extracted, settings.CHAT_MAX_EXTRACTED_CHARS)
        text_bytes = extracted.encode("utf-8")
        text_name = os.path.splitext(safe_name)[0] + ".txt"
        block = {"type": "document", "source": {"type": "file", "file_id": ""}}
        return text_name, "text/plain", text_bytes, block

    if ext in TEXT_EXTENSIONS and not data.strip():
        raise HTTPException(status_code=400, detail="No readable text found in file.")

    if ext not in DIRECT_MIME:
        raise HTTPException(status_code=400, detail=f"Unhandled extension: {ext}")

    block = {"type": "document", "source": {"type": "file", "file_id": ""}}
    return safe_name, DIRECT_MIME[ext], data, block


def _extract_office_text(ext: str, data: bytes) -> str:
    if ext == ".docx": return _parse_docx(data)
    if ext == ".xlsx": return _parse_xlsx(data)
    if ext == ".pptx": return _parse_pptx(data)
    raise ValueError(f"Unhandled office extension: {ext}")


def _parse_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables):
        parts.append(f"\n--- Table {i + 1} ---")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for name in wb.sheetnames:
        out.append(f"\n--- Sheet: {name} ---")
        for i, row in enumerate(wb[name].iter_rows(values_only=True)):
            if i >= 200:
                out.append("[Sheet truncated at 200 rows.]")
                break
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides):
        out.append(f"\n--- Slide {i + 1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                out.append(shape.text)
    return "\n".join(out)


def _validate_image(data: bytes) -> None:
    try:
        img = Image.open(io.BytesIO(data))
        img.verify()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid or unsafe image: {exc}")


async def upload_to_anthropic(client, filename: str, mime_type: str, data: bytes) -> str:
    file_obj = await client.beta.files.upload(
        file=(filename, data, mime_type),
        extra_headers={"anthropic-beta": "files-api-2025-04-14"},
    )
    return file_obj.id


async def delete_from_anthropic(client, file_id: str) -> None:
    try:
        await client.beta.files.delete(
            file_id,
            extra_headers={"anthropic-beta": "files-api-2025-04-14"},
        )
    except Exception:
        pass
